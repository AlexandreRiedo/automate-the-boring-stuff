from typing import cast

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt
from docx.styles.style import CharacterStyle, ParagraphStyle

# Document Setup --- Creating styles
doc = docx.Document()

st_para = cast(
    ParagraphStyle, doc.styles.add_style("Paragraph", WD_STYLE_TYPE.PARAGRAPH)
)
st_para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
st_para.paragraph_format.line_spacing = 1.5

st_script = cast(
    CharacterStyle, doc.styles.add_style("Script", WD_STYLE_TYPE.CHARACTER)
)
st_script.base_style = doc.styles["Normal"]
st_script.font.name = "Segoe Script"
st_script.font.size = Pt(16)

st_sans = cast(CharacterStyle, doc.styles.add_style("Sans", WD_STYLE_TYPE.CHARACTER))
st_sans.base_style = doc.styles["Normal"]
st_sans.font.name = "Segoe UI"
st_sans.font.size = Pt(16)

# Getting the guests
with open("guests.txt", "r") as file:
    guests = [f.strip() for f in file]

# Adding content to the Document
for idx, guest in enumerate(guests):
    para = doc.add_paragraph(style=st_para)
    para.add_run("It would be a pleasure to have the company of", st_script).add_break()
    para.add_run(guest, st_sans).bold = True
    para.runs[-1].add_break()
    para.add_run("at 11010 Memory Lane on the Evening of", st_script).add_break()
    para.add_run("April 1st", st_sans).add_break()
    para.add_run("at 7 o'clock", st_script)

    if idx < len(guests) - 1:
        para.runs[-1].add_break(WD_BREAK.PAGE)


# Writing the Document to file
doc.save("invitations.docx")
